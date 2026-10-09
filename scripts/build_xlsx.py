import sys, json, os
# 사용법: python3 build_xlsx.py <출력.xlsx> [데이터.json]
DATA = json.load(open(sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), 'data', 'complexes.json'), encoding='utf-8'))
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.worksheet.datavalidation import DataValidation

F='Arial'
BLUE=Font(name=F,color='0000FF',size=10); BLK=Font(name=F,size=10); GRN=Font(name=F,color='008000',size=10)
HDR=Font(name=F,bold=True,color='FFFFFF',size=10); TTL=Font(name=F,bold=True,size=14,color='1F3864')
HF=PatternFill('solid',fgColor='1F3864'); YEL=PatternFill('solid',fgColor='FFFF00'); INP=PatternFill('solid',fgColor='FFF2CC')
thin=Side(style='thin',color='A6A6A6'); BD=Border(left=thin,right=thin,top=thin,bottom=thin)
C=Alignment(horizontal='center',vertical='center',wrap_text=True); L=Alignment(horizontal='left',vertical='center',wrap_text=True)
NUM='#,##0.00;(#,##0.00);-'; PCT='0.0%;(0.0%);-'; MUL='0.00"배"'

wb=Workbook()
# ---------------- 안내 ----------------
g=wb.active; g.title='안내'
g['A1']=f"부산 서부권 34평 7개 단지 갭투자·수익률 자동계산 시트 (데이터 기준 {DATA['asOf']})"; g['A1'].font=TTL
rows=[
('사용법',''),
('1','[가정] 시트의 노란 칸(취득세율·중개보수율·시나리오 선택)을 수정하세요.'),
('2','[단지분석] 시트의 파란 글씨·연노랑 칸(매매가·전세가·상승률·평가점수)만 입력하세요. 나머지는 모두 수식으로 자동 계산됩니다.'),
('3','[순위] 시트는 종합점수 기준 1~7위를 자동 정렬합니다(동점이면 위쪽 행이 우선).'),
('4','[대출시뮬] 시트에서 실거주 매수(2안)의 대출이자와 월 부담액을 계산합니다.'),
('색상 규칙',''),
('파란 글씨','직접 입력하는 값(하드코딩 입력값)'),
('검은 글씨','수식(수정 금지)'),
('초록 글씨','다른 시트를 참조하는 수식'),
('노란 채우기','핵심 가정. 사용자가 바꾸는 칸'),
('데이터 출처·한계',''),
('출처',f"{DATA['asOf']} 기준 웹검색에 노출된 실거래·시세(호갱노노·KB부동산·리치고 등). 국토부 원문(rt.molit.go.kr)은 작성 환경의 네트워크 정책으로 직접 확인하지 못함."),
('주의','"추정" 표시 값은 애널리스트 추정치임. 계약 전 국토부 실거래가로 반드시 교체하세요. 수익률은 세전(양도세·보유세·매도 중개비 차감 전)이며 투자 권유가 아님.'),
('단위','금액: 억원 / 비율: % (소수로 저장, 예 0.05 = 5%)'),
]
for i,(a,b) in enumerate(rows,start=3):
    g.cell(i,1,a).font=Font(name=F,bold=(b==''),size=10); g.cell(i,2,b).font=BLK; g.cell(i,2).alignment=L
g['A9'].font=BLUE; g['A10'].font=BLK; g['A11'].font=GRN; g['A12'].fill=YEL
g.column_dimensions['A'].width=18; g.column_dimensions['B'].width=110

# ---------------- 가정 ----------------
a=wb.create_sheet('가정')
a['A1']='공통 가정'; a['A1'].font=TTL
arows=[
('취득세+지방교육세율',DATA['assumptions']['acqTaxRate'],PCT,'1주택 또는 비조정지역(부산) 2주택, 6억 전후 34평(85㎡ 이하 농특세 비과세) 기준 약 1.1%. 3주택 이상은 8% 중과 → 0.08로 변경'),
('매수 중개보수율',DATA['assumptions']['brokerRate'],PCT,'공인중개사법 부산 조례 기준 매매 5억~9억 상한 0.4%'),
('보유기간(년)',DATA['assumptions']['holdYears'],'0','예상시세 시나리오의 기간. 상승률은 [단지분석]에서 "보유기간 전체 누적 상승률"로 입력'),
('적용 시나리오',DATA['assumptions']['scenario'],None,'보수 / 기본 / 낙관 중 선택 → [순위] 시트의 수익률 열에 반영'),
('소액 갭 기준(억원 미만)',DATA['assumptions']['smallGap'],NUM,'레버리지 평가 "소액 갭 ◎" 기준'),
('중간 갭 기준(억원 미만)',DATA['assumptions']['midGap'],NUM,'이 값 이상이면 "갭 큼 △"'),
]
a['A3'],a['B3'],a['C3']='항목','값','근거/설명'
for c in 'ABC': a[f'{c}3'].font=HDR; a[f'{c}3'].fill=HF; a[f'{c}3'].alignment=C
for i,(k,v,fmt,n) in enumerate(arows,start=4):
    a.cell(i,1,k).font=BLK; c=a.cell(i,2,v); c.font=BLUE; c.fill=YEL; c.alignment=C
    if fmt: c.number_format=fmt
    a.cell(i,3,n).font=BLK; a.cell(i,3).alignment=L
    for j in (1,2,3): a.cell(i,j).border=BD
dv=DataValidation(type='list',formula1='"보수,기본,낙관"',allow_blank=False); a.add_data_validation(dv); dv.add('B7')
a.column_dimensions['A'].width=26; a.column_dimensions['B'].width=12; a.column_dimensions['C'].width=90
ACQ="가정!$B$4"; BRK="가정!$B$5"; SCN="가정!$B$7"; G1="가정!$B$8"; G2="가정!$B$9"

# ---------------- 단지분석 ----------------
s=wb.create_sheet('단지분석')
s['A1']='단지별 갭투자 분석 (34평, 단위: 억원)'; s['A1'].font=TTL
s['A2']='파란 글씨·연노랑 칸만 입력 → 나머지 자동계산'; s['A2'].font=Font(name=F,italic=True,size=9,color='595959')
hdr=['단지명','지역','세대수','입주','매매가\n(입력)','매매가\n추정여부','전세가\n(입력)','전세가\n추정여부',
 '부대비용\n(취득세+중개)','총매수가','갭투자금','회수율\n(전세가율)','투자금 대비\n전세가(배)','레버리지 평가',
 '상승률\n보수','상승률\n기본','상승률\n낙관',
 '예상시세\n보수','예상시세\n기본','예상시세\n낙관','수익금\n보수','수익금\n기본','수익금\n낙관','수익률\n보수','수익률\n기본','수익률\n낙관',
 '상승기대\n(30)','입지·교통\n(20)','갭·전세가율\n(15)','상품성\n(15)','학군\n(10)','공급리스크\n(10)','종합점수','순위','가격 근거(출처)']
H=4
for j,h in enumerate(hdr,1):
    c=s.cell(H,j,h); c.font=HDR; c.fill=HF; c.alignment=C; c.border=BD
data=[(c['name'],c['gu'],c['households'],c['moveIn'],c['sale'],c['saleType'],c['jeonse'],c['jeonseType'],
       c['growth']['bear'],c['growth']['base'],c['growth']['bull'],*c['scores'],c['saleBasis']+' / '+c['jeonseBasis']) for c in DATA['complexes']]
r0=H+1; rN=H+len(data)
for i,d in enumerate(data):
    r=r0+i
    (nm,gu,hh,yr,sale,se,jn,je,gb,gm,gu2,*sc,src)=d
    vals={1:nm,2:gu,3:hh,4:yr,5:sale,6:se,7:jn,8:je,15:gb,16:gm,17:gu2,35:src}
    for k,v in enumerate(sc): vals[27+k]=v
    for col,v in vals.items():
        c=s.cell(r,col,v); c.font=BLUE if col not in (1,) else Font(name=F,bold=True,size=10)
        if col in (5,7,15,16,17)+tuple(range(27,33)): c.fill=INP
    f={
     9:f'=E{r}*({ACQ}+{BRK})',10:f'=E{r}+I{r}',11:f'=E{r}-G{r}+I{r}',12:f'=IFERROR(G{r}/E{r},0)',13:f'=IFERROR(G{r}/K{r},0)',
     14:f'=IF(K{r}<{G1},"소액 갭 ◎",IF(K{r}<{G2},"중간 ○","갭 큼 △"))',
     18:f'=E{r}*(1+O{r})',19:f'=E{r}*(1+P{r})',20:f'=E{r}*(1+Q{r})',
     21:f'=R{r}-E{r}-I{r}',22:f'=S{r}-E{r}-I{r}',23:f'=T{r}-E{r}-I{r}',
     24:f'=IFERROR(U{r}/K{r},0)',25:f'=IFERROR(V{r}/K{r},0)',26:f'=IFERROR(W{r}/K{r},0)',
     33:f'=SUM(AA{r}:AF{r})',34:f'=COUNTIF($AG${r0}:$AG${rN},">"&AG{r})+COUNTIF($AG${r0}:AG{r},AG{r})',
    }
    for col,v in f.items():
        c=s.cell(r,col,v); c.font=BLK
    for col in range(1,36):
        c=s.cell(r,col); c.border=BD; c.alignment=L if col in (1,2,35) else C
    for col in (5,7,9,10,11,18,19,20,21,22,23): s.cell(r,col).number_format=NUM
    for col in (12,15,16,17,24,25,26): s.cell(r,col).number_format=PCT
    s.cell(r,13).number_format=MUL
# 평균 행
ra=rN+1
s.cell(ra,1,'평균').font=Font(name=F,bold=True,size=10)
for col in (5,7,9,10,11,12,13,18,19,20,21,22,23,24,25,26,33):
    L_=s.cell(4,col).column_letter
    c=s.cell(ra,col,f'=AVERAGE({L_}{r0}:{L_}{rN})'); c.font=BLK; c.number_format=s.cell(r0,col).number_format; c.alignment=C
for col in range(1,36): s.cell(ra,col).border=BD; s.cell(ra,col).fill=PatternFill('solid',fgColor='DCE6F2')
s.cell(ra+2,1,'주: 수익금 = 예상시세 − 매매가 − 부대비용(세전). 상승률·점수는 애널리스트 추정 가정이며 자유롭게 수정 가능. 점수 가중치: 상승기대30·입지20·갭15·상품성15·학군10·공급리스크10.').font=Font(name=F,italic=True,size=9)
s.conditional_formatting.add(f'Y{r0}:Y{rN}',ColorScaleRule(start_type='min',start_color='F8696B',mid_type='percentile',mid_value=50,mid_color='FFEB84',end_type='max',end_color='63BE7B'))
s.conditional_formatting.add(f'AH{r0}:AH{rN}',CellIsRule(operator='lessThanOrEqual',formula=['4'],fill=PatternFill('solid',fgColor='FFD966')))
widths={1:22,2:15,3:7,4:10,35:55}
for col in range(1,36): s.column_dimensions[s.cell(4,col).column_letter].width=widths.get(col,9.5)
s.row_dimensions[4].height=42; s.freeze_panes='B5'
s['E4'].comment=Comment('최신 국토부 실거래가로 교체하세요(억원).','Analyst')
s['O4'].comment=Comment('보유기간(가정!B6) 전체 누적 상승률. 애널리스트 추정.','Analyst')

# ---------------- 순위 ----------------
k=wb.create_sheet('순위')
k['A1']='종합점수 순위 (자동 정렬)'; k['A1'].font=TTL
k['A2']='="적용 시나리오: "&가정!$B$7'; k['A2'].font=GRN
kh=['순위','단지명','지역','종합점수','매매가','전세가','갭투자금','회수율','예상시세(선택 시나리오)','수익금(선택)','수익률(선택)','레버리지 평가']
for j,h in enumerate(kh,1):
    c=k.cell(4,j,h); c.font=HDR; c.fill=HF; c.alignment=C; c.border=BD
rng=lambda col: f"단지분석!${col}${r0}:${col}${rN}"
for i in range(1,8):
    r=4+i
    m=f"MATCH($A{r},{rng('AH')},0)"
    k.cell(r,1,i).font=BLK
    cols={2:'A',3:'B',4:'AG',5:'E',6:'G',7:'K',8:'L',12:'N'}
    for col,src in cols.items():
        k.cell(r,col,f"=INDEX({rng(src)},{m})").font=GRN
    k.cell(r,9,f'=IF({SCN}="보수",INDEX({rng("R")},{m}),IF({SCN}="낙관",INDEX({rng("T")},{m}),INDEX({rng("S")},{m})))').font=GRN
    k.cell(r,10,f'=IF({SCN}="보수",INDEX({rng("U")},{m}),IF({SCN}="낙관",INDEX({rng("W")},{m}),INDEX({rng("V")},{m})))').font=GRN
    k.cell(r,11,f'=IFERROR(J{r}/G{r},0)').font=BLK
    for col in range(1,13):
        c=k.cell(r,col); c.border=BD; c.alignment=C if col!=2 else L
        if i<=4: c.fill=PatternFill('solid',fgColor='FFF2CC')
    for col in (5,6,7,9,10): k.cell(r,col).number_format=NUM
    for col in (8,11): k.cell(r,col).number_format=PCT
for col,w in zip('ABCDEFGHIJKL',[6,24,16,9,9,9,10,9,14,11,11,12]): k.column_dimensions[col].width=w
k['A13']='1~4위(노란색) = 향후 가격 인상 기대 상위. 점수·가격 입력은 [단지분석]에서 수정.'; k['A13'].font=Font(name=F,italic=True,size=9)

# ---------------- 대출시뮬 ----------------
l=wb.create_sheet('대출시뮬')
l['A1']='실거주 매수 대출 시뮬레이션 (2안)'; l['A1'].font=TTL
l['A3'],l['B3']='항목','값'
for c in 'AB': l[f'{c}3'].font=HDR; l[f'{c}3'].fill=HF
lrows=[
('대상 단지(단지분석 행 선택)','대신2차푸르지오','text',True,'단지명을 목록에서 선택'),
('매매가(억원)',f'=INDEX(단지분석!$E${r0}:$E${rN},MATCH(B4,단지분석!$A${r0}:$A${rN},0))',NUM,False,''),
('부대비용(억원)',f'=INDEX(단지분석!$I${r0}:$I${rN},MATCH(B4,단지분석!$A${r0}:$A${rN},0))',NUM,False,''),
('LTV(담보인정비율)',0.7,PCT,True,'비규제지역 일반 70% 가정. 실제 한도는 DSR·소득에 따라 달라짐'),
('대출금(억원)','=B5*B7',NUM,False,''),
('필요 자기자본(억원)','=B5+B6-B8',NUM,False,''),
('대출금리(연)',0.045,PCT,True,'가정값. 기준금리 3.00%(2026.8.27) 환경의 주담대 금리로 은행 견적으로 교체'),
('대출기간(년)',30,'0',True,''),
('월 원리금(만원, 원리금균등)','=IFERROR(PMT(B10/12,B11*12,-B8*10000),0)','#,##0',False,'억원×10,000 = 만원'),
('월 이자만(만원, 첫달)','=B8*10000*B10/12','#,##0',False,''),
('보유기간 예상시세(기본, 억원)',f'=INDEX(단지분석!$S${r0}:$S${rN},MATCH(B4,단지분석!$A${r0}:$A${rN},0))',NUM,False,''),
('보유기간 총이자(억원, 근사)',f'=B8*B10*가정!$B$6',NUM,False,'원금 상환분 무시한 단순 근사'),
('실거주 순수익(억원, 세전)','=B14-B5-B6-B15',NUM,False,'1주택 2년 거주 시 양도세 비과세(12억 이하) 가정'),
('자기자본 수익률','=IFERROR(B16/B9,0)',PCT,False,''),
]
for i,(nm,v,fmt,inp,note) in enumerate(lrows,start=4):
    l.cell(i,1,nm).font=BLK; c=l.cell(i,2,v)
    c.font=BLUE if inp else (GRN if '단지분석' in str(v) or '가정' in str(v) else BLK)
    if inp: c.fill=YEL
    if fmt!='text': c.number_format=fmt
    c.alignment=C; l.cell(i,3,note).font=Font(name=F,size=9,color='595959')
    for j in (1,2): l.cell(i,j).border=BD
dv2=DataValidation(type='list',formula1=f'=단지분석!$A${r0}:$A${rN}'); l.add_data_validation(dv2); dv2.add('B4')
l.column_dimensions['A'].width=32; l.column_dimensions['B'].width=18; l.column_dimensions['C'].width=70

wb.save(sys.argv[1]); print('saved')
